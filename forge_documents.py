"""Bounded project-document extraction and evidence-backed artifact inspection."""
from __future__ import annotations

import csv
import base64
import hashlib
import io
import json
import os
from pathlib import Path
import re
import stat
import tempfile
import threading
import zipfile
from xml.etree import ElementTree

from project_tools import ProjectTools

MAX_DOCUMENT_BYTES = 32 * 1024 * 1024
MAX_ARCHIVE_BYTES = 64 * 1024 * 1024
MAX_EXTRACT_CHARACTERS = 48000
TEXT_SUFFIXES = {'.txt', '.md', '.json', '.html', '.css', '.js', '.ts', '.tsx', '.py', '.yaml', '.yml', '.xml', '.log', '.csv', '.tsv'}


def project_file(project, home, value):
    path = ProjectTools(project['path'], Path(home) / 'backups')._path(value)
    if not stat.S_ISREG(path.stat().st_mode) or path.stat().st_nlink > 1:
        raise ValueError('Select a regular project file.')
    if path.stat().st_size > MAX_DOCUMENT_BYTES:
        raise ValueError('Document exceeds the 32 MiB inspection limit.')
    return path


def _archive(path):
    package = zipfile.ZipFile(path)
    entries = package.infolist()
    if len(entries) > 10000 or sum(e.file_size for e in entries) > MAX_ARCHIVE_BYTES:
        package.close()
        raise ValueError('Document archive expands beyond the inspection limit.')
    if any(e.flag_bits & 1 for e in entries):
        package.close()
        raise ValueError('Encrypted documents are unsupported.')
    return package


def extract_document(path, start=0, limit=12000):
    """No macros, formulas, embedded objects or repository code are executed."""
    path = Path(path)
    if type(start) is not int or start < 0 or type(limit) is not int or not 1 <= limit <= MAX_EXTRACT_CHARACTERS:
        raise ValueError('Use a nonnegative offset and a limit from 1 to 48000.')
    if path.stat().st_size > MAX_DOCUMENT_BYTES:
        raise ValueError('Document exceeds the 32 MiB inspection limit.')
    suffix = path.suffix.lower()
    metadata = {'format': suffix.lstrip('.'), 'warnings': []}
    chunks = []
    bound = min(start + limit + 1, 2_000_000)
    if start > 2_000_000:
        raise ValueError('Extraction offset exceeds the 2M-character safety limit.')
    def append(value):
        if sum(len(part) for part in chunks) < bound:
            chunks.append(str(value)[:bound])
    if suffix in TEXT_SUFFIXES:
        text = path.read_bytes().decode('utf-8-sig')
        if suffix in ('.csv', '.tsv'):
            reader = csv.reader(io.StringIO(text), delimiter='\t' if suffix == '.tsv' else ',')
            rows = 0
            for row in reader:
                rows += 1
                append(json.dumps(row, ensure_ascii=False) + '\n')
                if sum(map(len, chunks)) >= bound:
                    break
            metadata['rows_inspected'] = rows
        else:
            append(text)
    elif suffix == '.docx':
        with _archive(path) as package:
            root = ElementTree.fromstring(package.read('word/document.xml'))
            namespace = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
            for paragraph in root.findall('.//w:p', namespace):
                append(''.join(node.text or '' for node in paragraph.findall('.//w:t', namespace)) + '\n')
                if sum(map(len, chunks)) >= bound:
                    break
        metadata['warnings'].append('Text and table paragraphs extracted; layout, comments and embedded media are not rendered.')
    elif suffix == '.xlsx':
        # Validate ZIP expansion before handing it to the optional reader.
        with _archive(path):
            pass
        try:
            import openpyxl
        except ImportError:
            raise ValueError('XLSX inspection needs openpyxl in the Forge runtime.') from None
        book = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            metadata['sheets'] = book.sheetnames[:100]
            for sheet in book.worksheets[:100]:
                append('Sheet: ' + sheet.title + '\n')
                for row in sheet.iter_rows(values_only=True):
                    append(json.dumps([str(value) if value is not None else '' for value in row[:200]], ensure_ascii=False) + '\n')
                    if sum(map(len, chunks)) >= bound:
                        break
                if sum(map(len, chunks)) >= bound:
                    break
        finally:
            book.close()
        metadata['warnings'].append('Formulas are returned as text and never recalculated; layout is not rendered.')
    elif suffix == '.pdf':
        try:
            from pypdf import PdfReader
        except ImportError:
            raise ValueError('PDF inspection needs pypdf in the Forge runtime.') from None
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError('Encrypted PDFs are unsupported.')
        metadata['pages'] = len(reader.pages)
        for index, page in enumerate(reader.pages[:200]):
            append(f'Page {index + 1}\n' + (page.extract_text() or '') + '\n')
            if sum(map(len, chunks)) >= bound:
                break
        metadata['warnings'].append('Text extraction only. Scanned pages may need OCR; visual layout is not verified.')
    else:
        raise ValueError('Supported formats: PDF, DOCX, XLSX, CSV, TSV and UTF-8 text.')
    text = ''.join(chunks)
    return {**metadata, 'text': text[start:start + limit], 'start': start,
            'next_start': start + limit if len(text) > start + limit else None,
            'truncated': len(text) > start + limit}


class DocumentManager:
    def __init__(self, store):
        self.store = store
        self.hash_cache = {}
        self.hash_lock = threading.RLock()

    def file_hash(self, path, refresh=False):
        path = Path(path)
        before = path.stat()
        key = (str(path.resolve()), before.st_mtime_ns, before.st_ctime_ns, before.st_size, before.st_ino)
        if not stat.S_ISREG(before.st_mode) or before.st_nlink > 1 or before.st_size > MAX_DOCUMENT_BYTES:
            raise ValueError('Artifact is not a bounded regular file.')
        with self.hash_lock:
            if not refresh and key in self.hash_cache:
                return self.hash_cache[key]
        digest = hashlib.sha256()
        count = 0
        descriptor = os.open(path, os.O_RDONLY | getattr(os, 'O_BINARY', 0) | getattr(os, 'O_NOFOLLOW', 0))
        with os.fdopen(descriptor, 'rb') as source:
            opened = os.fstat(source.fileno())
            if not stat.S_ISREG(opened.st_mode) or opened.st_nlink > 1:
                raise ValueError('Artifact changed or is linked.')
            while chunk := source.read(1024 * 1024):
                count += len(chunk)
                if count > MAX_DOCUMENT_BYTES:
                    raise ValueError('Artifact exceeds 32 MiB.')
                digest.update(chunk)
            after = os.fstat(source.fileno())
        current = path.stat()
        signature = lambda value: (value.st_mtime_ns, value.st_ctime_ns, value.st_size, value.st_ino)
        identity = lambda value: (value.st_mtime_ns, value.st_size, value.st_ino)
        # Windows path stat and descriptor stat can report different ctime
        # semantics. Compare like sources, while matching inode/size/mtime.
        if signature(before) != signature(current) or signature(opened) != signature(after) or identity(before) != identity(opened):
            raise ValueError('Artifact changed during inspection.')
        value = digest.hexdigest()
        with self.hash_lock:
            self.hash_cache = {k: v for k, v in self.hash_cache.items() if k[0] != key[0]}
            if len(self.hash_cache) >= 512:
                self.hash_cache.clear()
            self.hash_cache[key] = value
        return value

    def get_artifact(self, identifier):
        record = self.store.entity('builder_artifacts', identifier)
        try:
            project = self.store.get_project(record['project_id'])
            path = project_file(project, self.store.home, record['path'])
            if self.file_hash(path) != record['sha256']:
                raise ValueError('Artifact changed. Verify its current contents again.')
            return {**record, 'current_verified': True}
        except (ValueError, OSError) as exc:
            return {**record, 'status': 'stale', 'current_verified': False, 'current_error': str(exc)[:500]}

    def extract(self, data):
        if data.get('document_id'):
            document = self.store.entity('documents', data['document_id'])
            if not isinstance(document.get('filename'), str) or not re.fullmatch(r'[a-f0-9]{64}\.[a-z0-9]+', document['filename']):
                raise ValueError('Attached document metadata is invalid.')
            path = self.store.home / 'attachments' / 'documents' / document['filename']
            if (not path.is_file() or path.is_symlink() or path.stat().st_nlink > 1 or
                    path.stat().st_size > MAX_DOCUMENT_BYTES or not path.resolve().is_relative_to(self.store.home.resolve()) or
                    hashlib.sha256(path.read_bytes()).hexdigest() != document['sha256']):
                raise ValueError('Attached document changed or is unavailable.')
            display = document['name']
        else:
            if not data.get('path') or not data.get('project_id'):
                raise ValueError('Provide an attached document_id or a project_id and path.')
            project = self.store.get_project(data['project_id'])
            path = project_file(project, self.store.home, data['path'])
            display = str(path.relative_to(Path(project['path']).resolve()))
        result = extract_document(path, data.get('start', 0), data.get('limit', 12000))
        return {'ok': True, 'path': display,
                'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), **result}

    def upload(self, data):
        name = data.get('name')
        if not isinstance(name, str) or not name or len(name) > 240 or any(char in name for char in '/\\\0'):
            raise ValueError('Use a plain document filename of at most 240 characters.')
        suffix = Path(name).suffix.lower()
        if suffix not in TEXT_SUFFIXES | {'.pdf', '.docx', '.xlsx'}:
            raise ValueError('Supported attachments: PDF, DOCX, XLSX, CSV, TSV and UTF-8 text.')
        encoded = data.get('content_base64')
        if not isinstance(encoded, str) or len(encoded) > (MAX_DOCUMENT_BYTES + 2) // 3 * 4:
            raise ValueError('Document exceeds the 32 MiB attachment limit.')
        try:
            raw = base64.b64decode(encoded, validate=True)
        except (ValueError, __import__('binascii').Error):
            raise ValueError('Document content must be valid base64.') from None
        if not raw or len(raw) > MAX_DOCUMENT_BYTES:
            raise ValueError('Upload a nonempty document of at most 32 MiB.')
        if data.get('project_id'):
            self.store.get_project(data['project_id'])
        digest = hashlib.sha256(raw).hexdigest()
        directory = self.store.home / 'attachments' / 'documents'
        directory.mkdir(parents=True, exist_ok=True)
        if directory.is_symlink() or not directory.resolve().is_relative_to(self.store.home.resolve()):
            raise ValueError('Document attachment directory is unavailable.')
        filename = digest + suffix
        path = directory / filename
        if path.is_symlink():
            raise ValueError('Document attachment destination is unavailable.')
        if not path.exists():
            handle, temporary = tempfile.mkstemp(dir=directory, prefix='upload-')
            try:
                with os.fdopen(handle, 'wb') as output:
                    output.write(raw)
                os.replace(temporary, path)
            finally:
                if os.path.exists(temporary):
                    os.unlink(temporary)
        # Include the suffix in the ID so the same bytes cannot be reinterpreted
        # as a different format via an existing content-addressed reference.
        identifier = hashlib.sha256((digest + suffix).encode()).hexdigest()
        return self.store.save_entity('documents', {'id': identifier, 'name': name, 'filename': filename,
            'sha256': digest, 'bytes': len(raw), 'format': suffix.lstrip('.'), 'project_id': data.get('project_id'),
            'untrusted': True})

    def list(self, data):
        return {'documents': [{key: value for key, value in item.items() if key != 'filename'}
                             for item in self.store.entities('documents')
                             if not data.get('project_id') or item.get('project_id') == data['project_id']][-100:][::-1]}

    def create(self, data):
        project = self.store.get_project(data['project_id'])
        if data.get('builder_id') and self.store.entity('builders', data['builder_id'])['project_id'] != project['id']:
            raise ValueError('This artifact Builder belongs to another project.')
        tools = ProjectTools(project['path'], self.store.home / 'backups')
        path = tools._path(data['path'])
        kind = path.suffix.lower().lstrip('.')
        if kind not in ('txt', 'md', 'csv', 'tsv', 'xlsx', 'docx', 'pdf'):
            raise ValueError('Export formats: TXT, Markdown, CSV, TSV, XLSX, DOCX and PDF.')
        title, text = data.get('title', 'Export'), data.get('text', '')
        if not isinstance(title, str) or len(title) > 200 or not isinstance(text, str) or len(text) > 128000:
            raise ValueError('Use a title up to 200 and body up to 128000 characters.')
        rows = data.get('rows', [])
        if not isinstance(rows, list) or len(rows) > 5000 or any(not isinstance(row, list) or len(row) > 200 for row in rows):
            raise ValueError('Export tables support at most 5000 rows and 200 columns.')
        if any(value is not None and type(value) not in (str, int, float, bool) for row in rows for value in row):
            raise ValueError('Export cells must contain text, numbers, booleans or null.')
        if len(json.dumps(rows, ensure_ascii=False, allow_nan=False)) > 1_000_000:
            raise ValueError('Export table is too large.')
        buffer = io.BytesIO()
        if kind in ('txt', 'md'):
            raw = text.encode('utf-8')
        elif kind in ('csv', 'tsv'):
            output = io.StringIO(newline='')
            writer = csv.writer(output, delimiter='\t' if kind == 'tsv' else ',')
            writer.writerows(rows)
            raw = output.getvalue().encode('utf-8-sig')
        elif kind == 'xlsx':
            try:
                from openpyxl import Workbook
            except ImportError:
                raise ValueError('XLSX creation needs openpyxl in the Forge runtime.') from None
            book = Workbook()
            sheet = book.active
            sheet.title = 'Export'
            for row in rows:
                sheet.append(row)
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'  # Table values are data, not executable formulas.
            sheet.freeze_panes = 'A2' if len(rows) > 1 else None
            if rows:
                sheet.auto_filter.ref = sheet.dimensions
            book.save(buffer)
            raw = buffer.getvalue()
        elif kind == 'docx':
            try:
                from docx import Document
            except ImportError:
                raise ValueError('DOCX creation needs python-docx in the Forge runtime.') from None
            document = Document()
            document.add_heading(title, 0)
            for paragraph in text.splitlines():
                document.add_paragraph(paragraph)
            if rows:
                columns = max(map(len, rows))
                if columns:
                    table = document.add_table(rows=0, cols=columns)
                    table.style = 'Table Grid'
                    for row in rows:
                        cells = table.add_row().cells
                        for index, value in enumerate(row):
                            cells[index].text = str(value) if value is not None else ''
            document.save(buffer)
            raw = buffer.getvalue()
        else:
            try:
                from reportlab.lib.styles import getSampleStyleSheet
                from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
                from reportlab.lib import colors
            except ImportError:
                raise ValueError('PDF creation needs reportlab in the Forge runtime.') from None
            from xml.sax.saxutils import escape
            styles = getSampleStyleSheet()
            elements = [Paragraph(escape(title), styles['Title']), Spacer(1, 12)]
            elements.extend(Paragraph(escape(paragraph) or '&#160;', styles['BodyText']) for paragraph in text.splitlines())
            if rows:
                columns = max(map(len, rows))
                if columns > 12:
                    raise ValueError('PDF tables support at most 12 columns; use XLSX for wider tables.')
                if columns:
                    cells = [[Paragraph(escape(str(value)) if value is not None else '', styles['BodyText']) for value in row] +
                             [''] * (columns - len(row)) for row in rows]
                    table = Table(cells, repeatRows=1)
                    table.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), .5, colors.lightgrey),
                                               ('BACKGROUND', (0, 0), (-1, 0), colors.whitesmoke),
                                               ('VALIGN', (0, 0), (-1, -1), 'TOP')]))
                    elements.extend([Spacer(1, 12), table])
            SimpleDocTemplate(buffer).build(elements)
            raw = buffer.getvalue()
        if not raw:
            raise ValueError('Provide nonempty document text or table rows.')
        if len(raw) > MAX_DOCUMENT_BYTES:
            raise ValueError('Generated document exceeds 32 MiB.')
        # The same project boundary, hash conflict check, undo backup and atomic
        # replace are used for text and binary exports.
        if not path.parent.is_dir():
            tools.make_directory(str(path.parent))
        saved = tools.write_bytes(str(path), raw, data.get('expected_sha256', 'missing'))
        with self.hash_lock:
            self.hash_cache = {k: v for k, v in self.hash_cache.items() if k[0] != str(path.resolve())}
        verified = self.verify({'project_id': project['id'], 'builder_id': data.get('builder_id'),
                                'path': data['path'], 'expected_sha256': saved['sha256']})
        return {'ok': True, 'file': saved, 'verification': verified,
                'next_action': 'Inspect rendered pages or the native document before claiming visual quality.'}

    def download(self, data):
        record = self.store.entity('builder_artifacts', data['id'])
        project = self.store.get_project(record['project_id'])
        path = project_file(project, self.store.home, record['path'])
        raw = path.read_bytes()
        if hashlib.sha256(raw).hexdigest() != record['sha256']:
            raise ValueError('Artifact changed. Verify it again before download.')
        return {'name': path.name, 'content_base64': base64.b64encode(raw).decode(),
                'sha256': record['sha256'], 'bytes': len(raw), 'format': record['format']}

    def artifacts(self, data):
        items = self.store.entities('builder_artifacts')
        return {'artifacts': [self.get_artifact(item['id']) for item in items if
            (not data.get('project_id') or item['project_id'] == data['project_id']) and
            (not data.get('builder_id') or item.get('builder_id') == data['builder_id'])][-100:][::-1]}

    def verify(self, data):
        project = self.store.get_project(data['project_id'])
        if data.get('builder_id') and self.store.entity('builders', data['builder_id'])['project_id'] != project['id']:
            raise ValueError('This artifact Builder belongs to another project.')
        path = project_file(project, self.store.home, data['path'])
        digest = self.file_hash(path)
        if data.get('expected_sha256') and data['expected_sha256'] != digest:
            raise ValueError('Artifact changed since the supplied SHA-256. Inspect it again.')
        result = extract_document(path, 0, 12000)
        if self.file_hash(path, refresh=True) != digest:
            raise ValueError('Artifact changed during structural inspection. Verify it again.')
        # Verification is deliberately structural. A model cannot turn this into
        # an assertion of factual correctness or visual quality.
        record = self.store.save_entity('builder_artifacts', {'project_id': project['id'],
            'builder_id': data.get('builder_id'), 'path': str(path.relative_to(Path(project['path']).resolve())),
            'sha256': digest, 'bytes': path.stat().st_size, 'format': result['format'],
            'status': 'structure_verified', 'visual_verified': False,
            'warnings': result['warnings'], 'excerpt': result['text'][:4000]})
        evidence = self.store.artifact({'type': 'document_verification', **record})
        return {'ok': True, **record, 'evidence': evidence}
